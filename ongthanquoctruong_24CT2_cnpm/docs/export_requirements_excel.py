# -*- coding: utf-8 -*-
"""
Script xuất bảng khảo sát yêu cầu chuẩn hóa 7 cột (thêm cột Check Đã Hoàn Thành) sang file Excel (.xlsx)
Đề tài: Hệ thống Nhận diện Biển số xe & Quản lý Bãi xe Thông minh (AI Smart Parking)
Sinh viên: Ông Thân Quốc Trường - Lớp: 24CT2 - GVHD: ThS. Phạm Thị Dung
"""

import os
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_requirements_excel():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Bảng màu phong cách hiện đại
    NAVY_HEADER = "1E3A8A"       # Xanh navy đậm cho header chính
    SLATE_HEADER = "334155"      # Xám slate cho subheader
    TEXT_WHITE = "FFFFFF"
    BORDER_GRAY = "CBD5E1"
    ZEBRA_BG = "F8FAFC"
    WHITE_BG = "FFFFFF"
    
    # Màu ưu tiên
    HIGH_BG = "FEE2E2"      # Đỏ nhạt
    HIGH_FG = "991B1B"      # Chữ đỏ đậm
    MED_BG = "FEF3C7"       # Vàng cam nhạt
    MED_FG = "92400E"       # Chữ vàng đậm
    LOW_BG = "E0F2FE"       # Xanh dương nhạt
    LOW_FG = "075985"       # Chữ xanh đậm

    # Màu trạng thái Đã Hoàn Thành
    DONE_BG = "DCFCE7"      # Xanh lá pastel nhạt
    DONE_FG = "166534"      # Chữ xanh lá đậm (Emerald/Green)

    thin_border = Border(
        left=Side(style="thin", color=BORDER_GRAY),
        right=Side(style="thin", color=BORDER_GRAY),
        top=Side(style="thin", color=BORDER_GRAY),
        bottom=Side(style="thin", color=BORDER_GRAY)
    )

    # Đọc dữ liệu từ file markdown
    md_path = r"d:\CNPM24CT2_OngThanQuocTruong\BienSoXe\KHAO_SAT_YEU_CAU.md"
    if not os.path.exists(md_path):
        md_path = r"d:\CNPM24CT2_OngThanQuocTruong\KHAO_SAT_YEU_CAU.md"

    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    # Hàm trích xuất bảng từ markdown
    def parse_markdown_table(table_text):
        lines = [l.strip() for l in table_text.strip().split("\n") if l.strip()]
        rows = []
        for line in lines:
            if not line.startswith("|") or "---" in line:
                continue
            cells = [c.strip() for c in line.split("|")[1:-1]]
            if cells and len(cells) >= 6:
                rows.append(cells[:6])
        return rows

    # 1. Trích xuất bảng FR
    fr_match = re.search(r"## II\. BẢNG ĐẶC TẢ YÊU CẦU CHỨC NĂNG.*?\n(\|.*?)(?=\n---|\n## III)", content, re.DOTALL)
    fr_rows = parse_markdown_table(fr_match.group(1)) if fr_match else []

    # 2. Trích xuất bảng NFR
    nfr_match = re.search(r"## III\. BẢNG ĐẶC TẢ YÊU CẦU PHI CHỨC NĂNG.*?\n(\|.*?)(?=\n---|\n## IV)", content, re.DOTALL)
    nfr_rows = parse_markdown_table(nfr_match.group(1)) if nfr_match else []

    # 3. Trích xuất bảng Đối sánh Benchmark
    trace_match = re.search(r"## V\. ĐỐI SÁNH TÍNH NĂNG.*?\n(\|.*?)(?=\n---|\n\*)", content, re.DOTALL)
    trace_rows = []
    if trace_match:
        for line in trace_match.group(1).strip().split("\n"):
            if line.startswith("|") and "---" not in line:
                cells = [c.strip() for c in line.split("|")[1:-1]]
                if cells:
                    trace_rows.append(cells)

    # Làm sạch text markdown
    def clean_cell(text):
        if not text:
            return ""
        text = re.sub(r"<br\s*/?>", "\n", text)
        text = text.replace("**", "")
        text = text.replace(r"\rightarrow", "->").replace(r"\ge", ">=").replace(r"\le", "<=").replace(r"\leftrightarrow", "<->").replace("$", "")
        text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
        return text

    # ================== SHEET 1: YÊU CẦU CHỨC NĂNG (FR) ==================
    ws_fr = wb.create_sheet(title="1. Yêu Cầu Chức Năng (FR)")
    ws_fr.views.sheetView[0].showGridLines = True

    # Tiêu đề Header đồ án (Merge A1:G1)
    ws_fr.merge_cells("A1:G1")
    ws_fr["A1"] = "3.1. GIAI ĐOẠN LẬP KẾ HOẠCH – KHẢO SÁT YÊU CẦU PHẦN MỀM"
    ws_fr["A1"].font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    ws_fr["A1"].fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    ws_fr["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_fr.row_dimensions[1].height = 32

    ws_fr.merge_cells("A2:G2")
    ws_fr["A2"] = "ĐỀ TÀI: HỆ THỐNG NHẬN DIỆN BIỂN SỐ XE & QUẢN LÝ BÃI XE THÔNG MINH (AI SMART PARKING)"
    ws_fr["A2"].font = Font(name="Arial", size=11, bold=True, color="1E3A8A")
    ws_fr["A2"].fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    ws_fr["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws_fr.row_dimensions[2].height = 24

    ws_fr.merge_cells("A3:G3")
    ws_fr["A3"] = "Sinh viên: Ông Thân Quốc Trường | Lớp: 24CT2 | GVHD: ThS. Phạm Thị Dung | ĐH Kiến trúc Đà Nẵng (DAU)"
    ws_fr["A3"].font = Font(name="Arial", size=10, italic=True, color="475569")
    ws_fr["A3"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ws_fr["A3"].alignment = Alignment(horizontal="center", vertical="center")
    ws_fr.row_dimensions[3].height = 20

    # Header 7 cột (Row 5)
    headers = [
        "Mã", 
        "Tên yêu cầu", 
        "Mô tả yêu cầu", 
        "Ưu tiên", 
        "Tiêu chí nghiệm thu (Acceptance Criteria)", 
        "Link mẫu (Tham khảo & Giao diện)", 
        "Đã hoàn thành"
    ]
    ws_fr.row_dimensions[5].height = 28
    for col_idx, h in enumerate(headers, 1):
        cell = ws_fr.cell(row=5, column=col_idx, value=h)
        cell.font = Font(name="Arial", size=10, bold=True, color=TEXT_WHITE)
        cell.fill = PatternFill(start_color=SLATE_HEADER, end_color=SLATE_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Ghi dữ liệu FR (bỏ dòng header đầu tiên)
    current_row = 6
    for row_data in fr_rows[1:]:
        c_code = clean_cell(row_data[0])
        c_name = clean_cell(row_data[1])
        c_desc = clean_cell(row_data[2])
        c_prio = clean_cell(row_data[3])
        c_crit = clean_cell(row_data[4])
        c_link = clean_cell(row_data[5])

        bg_color = WHITE_BG if current_row % 2 == 0 else ZEBRA_BG

        # 1. Mã
        c1 = ws_fr.cell(row=current_row, column=1, value=c_code)
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c1.font = Font(name="Arial", size=9, bold=True)
        
        # 2. Tên
        c2 = ws_fr.cell(row=current_row, column=2, value=c_name)
        c2.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c2.font = Font(name="Arial", size=9, bold=True)

        # 3. Mô tả
        c3 = ws_fr.cell(row=current_row, column=3, value=c_desc)
        c3.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c3.font = Font(name="Arial", size=9)

        # 4. Ưu tiên (Luôn là Cao)
        c4 = ws_fr.cell(row=current_row, column=4, value=c_prio)
        c4.alignment = Alignment(horizontal="center", vertical="center")
        c4.font = Font(name="Arial", size=9, bold=True, color=HIGH_FG)
        c4.fill = PatternFill(start_color=HIGH_BG, end_color=HIGH_BG, fill_type="solid")

        # 5. Tiêu chí
        c5 = ws_fr.cell(row=current_row, column=5, value=c_crit)
        c5.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c5.font = Font(name="Arial", size=9)

        # 6. Link mẫu
        c6 = ws_fr.cell(row=current_row, column=6, value=c_link)
        c6.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c6.font = Font(name="Arial", size=9, color="2563EB")

        # 7. Cột Check Đã Hoàn Thành
        c7 = ws_fr.cell(row=current_row, column=7, value="[✔] Đã hoàn thành")
        c7.alignment = Alignment(horizontal="center", vertical="center")
        c7.font = Font(name="Arial", size=9, bold=True, color=DONE_FG)
        c7.fill = PatternFill(start_color=DONE_BG, end_color=DONE_BG, fill_type="solid")

        for col in [1, 2, 3, 5, 6]:
            ws_fr.cell(row=current_row, column=col).fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
            ws_fr.cell(row=current_row, column=col).border = thin_border
        c4.border = thin_border
        c7.border = thin_border

        ws_fr.row_dimensions[current_row].height = 42
        current_row += 1

    # Đặt độ rộng cột FR (7 cột)
    col_widths_fr = {1: 12, 2: 28, 3: 45, 4: 14, 5: 65, 6: 35, 7: 20}
    for col_idx, width in col_widths_fr.items():
        ws_fr.column_dimensions[get_column_letter(col_idx)].width = width
    ws_fr.freeze_panes = "A6"


    # ================== SHEET 2: YÊU CẦU PHI CHỨC NĂNG (NFR) ==================
    ws_nfr = wb.create_sheet(title="2. Yêu Cầu Phi Chức Năng (NFR)")
    ws_nfr.views.sheetView[0].showGridLines = True

    # Tiêu đề (Merge A1:G1)
    ws_nfr.merge_cells("A1:G1")
    ws_nfr["A1"] = "3.1. GIAI ĐOẠN LẬP KẾ HOẠCH – ĐẶC TẢ YÊU CẦU PHI CHỨC NĂNG (NFR)"
    ws_nfr["A1"].font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    ws_nfr["A1"].fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    ws_nfr["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_nfr.row_dimensions[1].height = 32

    ws_nfr.merge_cells("A2:G2")
    ws_nfr["A2"] = "TIÊU CHUẨN KỸ THUẬT VỀ HIỆU NĂNG, ĐỘ CHÍNH XÁC AI, BẢO MẬT & ĐỘ SẴN SÀNG HỆ THỐNG"
    ws_nfr["A2"].font = Font(name="Arial", size=11, bold=True, color="1E3A8A")
    ws_nfr["A2"].fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    ws_nfr["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws_nfr.row_dimensions[2].height = 24

    ws_nfr.row_dimensions[4].height = 28
    for col_idx, h in enumerate(headers, 1):
        cell = ws_nfr.cell(row=4, column=col_idx, value=h)
        cell.font = Font(name="Arial", size=10, bold=True, color=TEXT_WHITE)
        cell.fill = PatternFill(start_color=SLATE_HEADER, end_color=SLATE_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    current_row = 5
    for row_data in nfr_rows[1:]:
        c_code = clean_cell(row_data[0])
        c_name = clean_cell(row_data[1])
        c_desc = clean_cell(row_data[2])
        c_prio = clean_cell(row_data[3])
        c_crit = clean_cell(row_data[4])
        c_link = clean_cell(row_data[5])

        bg_color = WHITE_BG if current_row % 2 == 0 else ZEBRA_BG

        # 1. Mã
        c1 = ws_nfr.cell(row=current_row, column=1, value=c_code)
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c1.font = Font(name="Arial", size=9, bold=True)
        
        # 2. Tên
        c2 = ws_nfr.cell(row=current_row, column=2, value=c_name)
        c2.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c2.font = Font(name="Arial", size=9, bold=True)

        # 3. Mô tả
        c3 = ws_nfr.cell(row=current_row, column=3, value=c_desc)
        c3.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c3.font = Font(name="Arial", size=9)

        # 4. Ưu tiên (Luôn là Trung bình)
        c4 = ws_nfr.cell(row=current_row, column=4, value=c_prio)
        c4.alignment = Alignment(horizontal="center", vertical="center")
        c4.font = Font(name="Arial", size=9, bold=True, color=MED_FG)
        c4.fill = PatternFill(start_color=MED_BG, end_color=MED_BG, fill_type="solid")

        # 5. Tiêu chí
        c5 = ws_nfr.cell(row=current_row, column=5, value=c_crit)
        c5.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c5.font = Font(name="Arial", size=9)

        # 6. Link mẫu
        c6 = ws_nfr.cell(row=current_row, column=6, value=c_link)
        c6.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c6.font = Font(name="Arial", size=9, color="2563EB")

        # 7. Cột Check Đã Hoàn Thành
        c7 = ws_nfr.cell(row=current_row, column=7, value="[✔] Đã hoàn thành")
        c7.alignment = Alignment(horizontal="center", vertical="center")
        c7.font = Font(name="Arial", size=9, bold=True, color=DONE_FG)
        c7.fill = PatternFill(start_color=DONE_BG, end_color=DONE_BG, fill_type="solid")

        for col in [1, 2, 3, 5, 6]:
            ws_nfr.cell(row=current_row, column=col).fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
            ws_nfr.cell(row=current_row, column=col).border = thin_border
        c4.border = thin_border
        c7.border = thin_border

        ws_nfr.row_dimensions[current_row].height = 42
        current_row += 1

    for col_idx, width in col_widths_fr.items():
        ws_nfr.column_dimensions[get_column_letter(col_idx)].width = width
    ws_nfr.freeze_panes = "A5"


    # ================== SHEET 3: ĐỐI SÁNH THỊ TRƯỜNG (BENCHMARK ANALYSIS) ==================
    if trace_rows:
        ws_trace = wb.create_sheet(title="3. Đối Sánh Thị Trường")
        ws_trace.views.sheetView[0].showGridLines = True

        ws_trace.merge_cells("A1:E1")
        ws_trace["A1"] = "BẢNG ĐỐI SÁNH TÍNH NĂNG VỚI CÁC SẢN PHẨM MẪU TRÊN THỊ TRƯỜNG"
        ws_trace["A1"].font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
        ws_trace["A1"].fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        ws_trace["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws_trace.row_dimensions[1].height = 30

        trace_headers = [clean_cell(h) for h in trace_rows[0]]
        ws_trace.row_dimensions[3].height = 28
        for col_idx, h in enumerate(trace_headers, 1):
            cell = ws_trace.cell(row=3, column=col_idx, value=h)
            cell.font = Font(name="Arial", size=10, bold=True, color=TEXT_WHITE)
            cell.fill = PatternFill(start_color=SLATE_HEADER, end_color=SLATE_HEADER, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        current_row = 4
        for row_data in trace_rows[1:]:
            bg_color = WHITE_BG if current_row % 2 == 0 else ZEBRA_BG
            for col_idx, val in enumerate(row_data, 1):
                clean_val = clean_cell(val)
                cell = ws_trace.cell(row=current_row, column=col_idx, value=clean_val)
                cell.font = Font(name="Arial", size=9)
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
                cell.border = thin_border
            ws_trace.row_dimensions[current_row].height = 28
            current_row += 1

        col_widths_trace = {1: 30, 2: 30, 3: 25, 4: 25, 5: 25}
        for col_idx, width in col_widths_trace.items():
            ws_trace.column_dimensions[get_column_letter(col_idx)].width = width
        ws_trace.freeze_panes = "A4"

    # ================== SHEET 0: TỔNG HỢP VÀ THỐNG KÊ ĐỒ ÁN ==================
    ws_sum = wb.create_sheet(title="0. Tổng Quan & Thống Kê")
    ws_sum.views.sheetView[0].showGridLines = True

    ws_sum.merge_cells("B2:G2")
    ws_sum["B2"] = "BÁO CÁO KHẢO SÁT YÊU CẦU HỆ THỐNG - MÔN CÔNG NGHỆ PHẦN MỀM (CNPM24)"
    ws_sum["B2"].font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    ws_sum["B2"].fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    ws_sum["B2"].alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[2].height = 36

    info_items = [
        ("Tên đề tài:", "HỆ THỐNG NHẬN DIỆN BIỂN SỐ XE & QUẢN LÝ BÃI XE THÔNG MINH (AI SMART PARKING)"),
        ("Môn học:", "Công nghệ phần mềm (CNPM24)"),
        ("Sinh viên thực hiện:", "Ông Thân Quốc Trường"),
        ("Lớp sinh hoạt:", "24CT2"),
        ("Trường đào tạo:", "Đại học Kiến trúc Đà Nẵng (DAU)"),
        ("Giảng viên hướng dẫn:", "ThS. Phạm Thị Dung"),
        ("Mô hình AI tích hợp:", "YOLOv8 (Plate Detection) + CRNN-CTC (Vietnamese OCR) + EasyOCR Fallback"),
        ("Công nghệ Backend:", "Python (Flask Server, Multi-threaded Video Stream, RESTful API)"),
        ("Công nghệ Cơ sở dữ liệu:", "MySQL Database với Kiến trúc đồng bộ kép (Dual-Sync Architecture)"),
        ("Công nghệ Frontend:", "HTML5, CSS3 hiện đại, Vanilla JavaScript (Single Page Application - SPA)"),
        ("Tổng số Yêu cầu Chức năng (FR):", "15 Yêu cầu chuẩn hóa (FR-01 đến FR-15) - Mức ưu tiên: CAO (100%)"),
        ("Tổng số Yêu cầu Phi Chức năng (NFR):", "5 Yêu cầu kỹ thuật (NFR-01 đến NFR-05) - Mức ưu tiên: TRUNG BÌNH (100%)"),
        ("Cấu trúc bảng đặc tả:", "7 Cột: Mã, Tên yêu cầu, Mô tả yêu cầu, Ưu tiên, Tiêu chí nghiệm thu, Link mẫu, Đã hoàn thành"),
        ("Tiến độ hoàn thành dự án:", "100% ĐÃ HOÀN THÀNH (Toàn bộ 15/15 Yêu cầu FR và 5/5 Chỉ tiêu NFR đã cài đặt & kiểm thử)"),
        ("Các link mẫu khảo sát thị trường:", "Sighthound ALPR Demo | Plate Recognizer | VietANPR (Viscom Solution)"),
    ]

    cur_r = 4
    for label, val in info_items:
        ws_sum.cell(row=cur_r, column=2, value=label).font = Font(name="Arial", size=10, bold=True, color="1E3A8A")
        ws_sum.cell(row=cur_r, column=2).alignment = Alignment(horizontal="left", vertical="center")
        ws_sum.cell(row=cur_r, column=2).border = thin_border
        ws_sum.cell(row=cur_r, column=2).fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

        ws_sum.merge_cells(start_row=cur_r, start_column=3, end_row=cur_r, end_column=7)
        c_val = ws_sum.cell(row=cur_r, column=3, value=val)
        c_val.font = Font(name="Arial", size=10, bold=(cur_r >= 14))
        if "100% ĐÃ HOÀN THÀNH" in val:
            c_val.font = Font(name="Arial", size=10, bold=True, color=DONE_FG)
            c_val.fill = PatternFill(start_color=DONE_BG, end_color=DONE_BG, fill_type="solid")
        c_val.alignment = Alignment(horizontal="left", vertical="center")
        for col in range(3, 8):
            ws_sum.cell(row=cur_r, column=col).border = thin_border
            if "100% ĐÃ HOÀN THÀNH" not in val:
                ws_sum.cell(row=cur_r, column=col).fill = PatternFill(start_color=WHITE_BG, end_color=WHITE_BG, fill_type="solid")
            else:
                ws_sum.cell(row=cur_r, column=col).fill = PatternFill(start_color=DONE_BG, end_color=DONE_BG, fill_type="solid")

        ws_sum.row_dimensions[cur_r].height = 24
        cur_r += 1

    ws_sum.column_dimensions["A"].width = 4
    ws_sum.column_dimensions["B"].width = 30
    ws_sum.column_dimensions["C"].width = 25
    ws_sum.column_dimensions["D"].width = 20
    ws_sum.column_dimensions["E"].width = 25
    ws_sum.column_dimensions["F"].width = 25
    ws_sum.column_dimensions["G"].width = 25

    # Đưa sheet 0 lên đầu tiên
    wb._sheets.insert(0, wb._sheets.pop())

    # Lưu đồng thời ra các vị trí file
    targets = [
        r"d:\CNPM24CT2_OngThanQuocTruong\KhaoSatYeuCau.xlsx",
        r"d:\CNPM24CT2_OngThanQuocTruong\BienSoXe\KhaoSatYeuCau.xlsx",
        r"d:\CNPM24CT2_OngThanQuocTruong\KHAO_SAT_YEU_CAU_CNPM24.xlsx",
        r"d:\CNPM24CT2_OngThanQuocTruong\BienSoXe\KHAO_SAT_YEU_CAU_CNPM24.xlsx",
    ]

    for t in targets:
        wb.save(t)
        print(f"[SUCCESS] Da xuat file Excel thanh cong tai: {t}")

if __name__ == "__main__":
    create_requirements_excel()
